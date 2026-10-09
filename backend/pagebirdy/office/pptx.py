"""PowerPoint (.pptx): the text on slides.

Text boxes, titles and other placeholders, text inside shapes, shapes inside
groups (at any depth) and table cells. Speaker notes, slide masters, layouts
and the theme are not read, and are written back byte for byte along with
media, charts, animations and every relationship (`package.write_patched`):
only a slide whose text changed is re-serialised.

Each `a:p` goes to the engine whole, its differently formatted runs tagged
(`office.runs`); a line break (`a:br`) is `⟦br⟧` and a field (`a:fld`, a slide
number or date) an opaque `⟦xN⟧`. A rewritten paragraph keeps its `a:pPr` and
`a:endParaRPr`; each span's text goes into a copy of the span's first run, so
its font, size, colour and weight come with it. Shape geometry is never
written. After a frame is rewritten, `pptx_fit` fits the new text inside it.

Pictures on slides (`p:pic`, in groups too) are read by OCR and redrawn in
place through `office.images`: one media part is one picture however many
slides show it, and it is written back in its own format at its own pixel
size, so the `p:pic` that places it is never touched. Pictures on layouts and
masters, picture fills on ordinary shapes, linked (not embedded) pictures and
SVG are not read.
"""

from __future__ import annotations

import copy
import re
import zipfile

import pptx as python_pptx
from lxml import etree
from pptx.shapes.group import GroupShape

from pagebirdy.office import images
from pagebirdy.office.adapter import Issue, blank_segment, reject_tags, writable
from pagebirdy.office.formats import MB, Format
from pagebirdy.office.package import member, part_bytes, write_patched
from pagebirdy.office.protect import restore
from pagebirdy.office.runs import Piece, decode, encode

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_SVG_BLIP = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_IGNORED_ATTRS = ("lang", "altLang", "dirty", "err", "smtClean", "smtId", "noProof")
_NEWLINES = re.compile(r"[\r\n  ]")


def _shapes(shapes):
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _shapes(shape.shapes)
        else:
            yield shape


def frames(prs):
    """(slide number, slide, shape, txBody, id prefix, is_table_cell) for every
    text body the adapter reads, in a fixed order."""
    for i, slide in enumerate(prs.slides, 1):
        # A shape is addressed by its position in tree order, not its `cNvPr`
        # id: copied or generated decks repeat ids within a slide. Ids carry no
        # "/" so they travel in a URL path segment (PATCH /api/segments/...).
        for k, shape in enumerate(_shapes(slide.shapes)):
            prefix = f"pptx:s{i}.sh{k}"
            if shape.has_text_frame:
                yield i, slide, shape, shape.text_frame._txBody, prefix, False
            elif getattr(shape, "has_table", False) and shape.has_table:
                for r, row in enumerate(shape.table.rows):
                    for c, cell in enumerate(row.cells):
                        yield i, slide, shape, cell._tc.txBody, f"{prefix}.c{r}-{c}", True


def _key(rpr) -> str:
    if rpr is None:
        return ""
    clean = copy.deepcopy(rpr)
    for attr in _IGNORED_ATTRS:
        clean.attrib.pop(attr, None)
    return etree.tostring(clean, method="c14n").decode("utf-8")


def _pieces(p) -> list[Piece]:
    pieces = []
    for child in p:
        if child.tag == A + "r":
            pieces.append(Piece("text", child.findtext(A + "t") or "", _key(child.find(A + "rPr")),
                                handle=child))
        elif child.tag == A + "br":
            pieces.append(Piece("break", handle=child))
        elif child.tag == A + "fld":
            pieces.append(Piece("object", handle=child))
    return pieces


def _text_runs(template, text: str) -> list:
    """`text` in the style of `template`, an `a:br` wherever it breaks."""
    out = []
    rpr = template.find(A + "rPr")
    for i, part in enumerate(_NEWLINES.split(text)):
        if i:
            br = etree.Element(A + "br")
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
            out.append(br)
        if part:
            r = copy.deepcopy(template)
            for child in list(r):
                if child.tag != A + "rPr":
                    r.remove(child)
            etree.SubElement(r, A + "t").text = part
            out.append(r)
    return out


def _rebuild_paragraph(p, layout, outs, placeholders) -> None:
    new = []
    breaks = iter(layout.breaks)
    base = layout.spans[layout.base][0].handle if layout.base >= 0 else None
    for o in outs:
        if o.kind == "text":
            new.extend(_text_runs(layout.spans[o.span][0].handle, restore(o.text, placeholders)))
        elif o.kind == "object":
            new.append(o.piece.handle)
        else:
            src = next(breaks, None)
            if src is not None:
                new.append(src.handle)
            elif base is not None:
                new.extend(_text_runs(base, "\n")[:1])
    ppr = p.find(A + "pPr")
    tail = [c for c in p if c.tag not in (A + "pPr", A + "r", A + "br", A + "fld")]
    for child in list(p):
        p.remove(child)
    for el in ([ppr] if ppr is not None else []) + new + tail:
        p.append(el)


def _has_text(p) -> bool:
    return any((r.findtext(A + "t") or "").strip() for r in p.findall(A + "r"))


def _picture_part(shape):
    """The embedded image part a `p:pic` shows, or None (not a picture, a
    linked picture, or an SVG whose embedded part is only its PNG fallback)."""
    el = shape._element
    if el.tag != P + "pic":
        return None
    blip = el.find(f"{P}blipFill/{A}blip")
    if blip is None or not blip.get(R + "embed"):
        return None
    if any(True for _ in blip.iter(_SVG_BLIP)):
        return None
    try:
        return shape.part.related_part(blip.get(R + "embed"))
    except KeyError:
        return None


def _collect(container_shapes, where: str, page: int, by_part: dict) -> None:
    for shape in _shapes(container_shapes):
        part = _picture_part(shape)
        if part is None:
            continue
        name = str(part.partname)
        m = by_part.get(name)
        if m is None:
            m = by_part[name] = images.Media(name, part.blob, part.content_type)
        m.where.append(where)
        m.pages.append(page)
        m.crops.append(tuple(float(getattr(shape, f"crop_{side}", 0.0) or 0.0)
                             for side in ("left", "top", "right", "bottom")))
        m.display_emu = (max(m.display_emu[0], int(shape.width or 0)),
                         max(m.display_emu[1], int(shape.height or 0)))


def pictures(prs) -> list[images.Media]:
    """Every picture part on a slide, once, in slide order."""
    by_part: dict = {}
    for i, slide in enumerate(prs.slides, 1):
        _collect(slide.shapes, f"slide {i}", i, by_part)
    return list(by_part.values())


def _template_pictures(prs) -> int:
    """Readable picture parts shown only on layouts and masters (never read)."""
    by_part: dict = {}
    for master in prs.slide_masters:
        _collect(master.shapes, "master", 0, by_part)
        for layout in master.slide_layouts:
            _collect(layout.shapes, "layout", 0, by_part)
    on_slides = {m.partname for m in pictures(prs)}
    return sum(1 for name, m in by_part.items()
               if name not in on_slides and images.eligible(m) is None)


def _readable(prs) -> tuple[list[images.Media], list[tuple[images.Media, str]]]:
    """(pictures to read, (picture, reason) for the rest). The first
    `images.max_images()` readable ones in slide order are read."""
    keep, skipped = [], []
    for m in pictures(prs):
        reason = images.eligible(m)
        if reason is None and len(keep) >= images.max_images():
            reason = f"over the limit of {images.max_images()} pictures per presentation"
        if reason:
            skipped.append((m, reason))
        else:
            keep.append(m)
    return keep, skipped


class PptxAdapter:
    def extract(self, src: str, *, allow_ocr: bool = True):
        """`allow_ocr=False` (a rebuild from review) reads pictures only from
        the OCR cache the first run wrote, never from the OCR service."""
        prs = python_pptx.Presentation(src)
        segs = []
        for slide_no, _, _, body, prefix, _ in frames(prs):
            for n, p in enumerate(body.findall(A + "p")):
                if _has_text(p):
                    tagged, _ = encode(_pieces(p))
                    segs.append(blank_segment(f"{prefix}.p{n}", tagged, page=slide_no))
        if allow_ocr and not images.enabled():
            return segs
        readable, _ = _readable(prs)
        regions, _ = images.load_regions(src, readable, allow_ocr=allow_ocr)
        for m in readable:
            if m.sha1 in regions:
                segs.extend(images.segments_for("pptx", m, regions[m.sha1]))
        return segs

    def rebuild(self, src, segments, out, lang):
        prs = python_pptx.Presentation(src)
        by_id = {s.id: s for s in segments}
        issues: list[Issue] = []
        touched_slides = {}
        written = []
        from pagebirdy.office import pptx_fit

        for slide_no, slide, shape, body, prefix, is_cell in list(frames(prs)):
            frame_written = False
            source_needed = None
            if not is_cell and any(writable(by_id.get(f"{prefix}.p{n}"))
                                   for n in range(len(body.findall(A + "p")))):
                source_needed = pptx_fit.needed_height(slide, shape, body, lang)
            for n, p in enumerate(body.findall(A + "p")):
                seg = by_id.get(f"{prefix}.p{n}")
                if not writable(seg):
                    continue
                _, layout = encode(_pieces(p))
                outs = decode(seg.target, layout)
                if outs is None:
                    issues.append(reject_tags(seg, f"slide {slide_no}"))
                    continue
                _rebuild_paragraph(p, layout, outs, seg.placeholders)
                written.append((p, pptx_fit.resolved_algn(slide, shape, p)))
                frame_written = True
            if frame_written:
                touched_slides[slide_no] = slide
                if not is_cell:
                    issue = pptx_fit.fit_frame(slide, shape, body, slide_no, lang, source_needed)
                    if issue is not None:
                        issues.append(issue)
        if lang.direction == "rtl" and written:
            from pagebirdy.office import rtl
            rtl.apply_pptx(written, lang)
        parts = {member(s.part.partname): part_bytes(s.part._element)
                 for s in touched_slides.values()}
        parts.update(self._pictures(src, prs, by_id, lang, issues))
        write_patched(src, out, parts)
        return issues

    def _pictures(self, src, prs, by_id, lang, issues) -> dict:
        """{media member: redrawn bytes} for every picture with a translation
        to draw; why any other picture kept its English goes to `issues`."""
        if not any(s.in_image for s in by_id.values()) and not images.enabled():
            return {}
        readable, skipped = _readable(prs)
        regions, _ = images.load_regions(src, readable, allow_ocr=False)
        out = {}
        for m in readable:
            if m.sha1 not in regions:
                issues.append(Issue("warning", "image", m.label,
                                    "picture was not read by OCR; its text is untranslated"))
                continue
            data, found = images.render("pptx", m, regions[m.sha1], by_id, lang)
            issues.extend(found)
            if data is not None:
                out[member(m.partname)] = data
        for m, reason in skipped:
            if "too small" not in reason:
                issues.append(Issue("warning", "image", m.label,
                                    f"picture not translated: {reason}"))
        n = _template_pictures(prs)
        if n:
            issues.append(Issue("warning", "image", "layouts/masters",
                                f"{n} picture(s) on slide layouts or masters were not translated"))
        return out

    def validate(self, src, out):
        try:
            a, b = python_pptx.Presentation(src), python_pptx.Presentation(out)
        except Exception as e:  # noqa: BLE001 — any failure to open is the finding
            return [Issue("error", "unreadable", "output", f"output does not open: {e}")]
        issues = []
        if len(a.slides) != len(b.slides):
            return [Issue("error", "structure", "presentation",
                          f"{len(a.slides)} slides in source, {len(b.slides)} in output")]
        if (a.slide_width, a.slide_height) != (b.slide_width, b.slide_height):
            issues.append(Issue("error", "structure", "presentation", "slide size changed"))
        for i, (sa, sb) in enumerate(zip(a.slides, b.slides), 1):
            if _geometry(sa) != _geometry(sb):
                issues.append(Issue("error", "geometry", f"slide {i}", "a shape moved or resized"))
            if "⟦" in etree.tostring(sb.part._element, encoding="unicode"):
                issues.append(Issue("error", "tokens", f"slide {i}",
                                    "a placeholder token reached the file"))
        issues += _validate_media(src, out)
        return issues

    def units(self, src):
        prs = python_pptx.Presentation(src)
        out = {"slides": len(prs.slides)}
        if images.enabled():
            readable, skipped = _readable(prs)
            out.update(images=len(readable), images_skipped=len(skipped))
        return out


def _validate_media(src, out) -> list[Issue]:
    """A redrawn picture must still be the same format at the same pixel size:
    the slide that places it was not rewritten to expect anything else."""
    issues = []
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(out) as b:
        names = set(b.namelist())
        for name in a.namelist():
            if not name.startswith("ppt/media/") or name not in names:
                continue
            old, new = a.read(name), b.read(name)
            if old == new:
                continue
            try:
                same = images.same_pixels_size(old, new)
            except Exception:  # noqa: BLE001 — an unreadable picture is the finding
                same = False
            if not same:
                issues.append(Issue("error", "image", name,
                                    "a redrawn picture changed format or size"))
    return issues


def _geometry(slide) -> list:
    out = []
    for el in slide.part._element.iter():
        if el.tag in (P + "cNvPr",):
            out.append(("id", el.get("id")))
        elif el.tag in (A + "off", A + "ext", A + "chOff", A + "chExt"):
            out.append((el.tag, el.get("x"), el.get("y"), el.get("cx"), el.get("cy")))
    return out


ADAPTER = PptxAdapter()
FORMAT = Format(".pptx", "pptx", "PPTX",
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                200 * MB, "ppt/presentation.xml", "pagebirdy.office.pptx")
