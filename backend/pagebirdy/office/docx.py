"""Word (.docx): paragraphs of the body, tables, text boxes, headers and footers.

python-docx opens and saves the package; the work is on the lxml tree, one
`w:p` at a time, in document order (`iter`), which reaches table cells (nested
tables too), text boxes (`w:txbxContent`, both the `mc:Choice` copy and its
`mc:Fallback` twin, translated identically because `Translator` dedups them)
and content controls. Footnotes, endnotes and comments are not read.

A paragraph's runs become `runs.Piece`s keyed by their formatting, so the
whole sentence goes to the engine with `⟦rN⟧` tags around differently styled
words. Anything that is not text is an opaque `⟦xN⟧` kept in place: a tab, a
drawing, a footnote reference, a whole field (simple or complex), a content
control, deleted text, an equation. Rebuilding a paragraph replaces only its
runs: each span's text goes into a copy of that span's first run (so its
`w:rPr` -- bold, colour, font, size -- comes with it), objects and line breaks
are moved back in where the translation put them, a hyperlink is re-created
around the words it now covers, and `w:pPr`, bookmarks and the like stay.
Styles, numbering, sections, images, page breaks are never touched.
"""

from __future__ import annotations

import copy
import re

import docx as python_docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from lxml import etree

from pagebirdy.office.adapter import Issue, blank_segment, reject_tags, writable
from pagebirdy.office.formats import MB, Format
from pagebirdy.office.package import member, part_bytes, write_patched
from pagebirdy.office.protect import restore
from pagebirdy.office.runs import Piece, decode, encode

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# Containers whose runs are part of the sentence, re-created around the words
# they cover after translation.
_WRAPPERS = {W + "hyperlink", W + "ins", W + "smartTag"}
# Elements kept whole and opaque: fields, controls, tracked deletions, math.
_WHOLE = {W + "fldSimple", W + "sdt", W + "del", W + "moveFrom", W + "moveTo",
          W + "customXml", M + "oMath", M + "oMathPara"}
# Formatting that says nothing about how the text looks.
_IGNORED_RPR = {W + "lang", W + "noProof", W + "proofErr"}
_NEWLINES = re.compile(r"[\r\n  ]")


def _key(rpr, wrapper_index: int) -> str:
    if rpr is None:
        body = ""
    else:
        clean = copy.deepcopy(rpr)
        for child in list(clean):
            if child.tag in _IGNORED_RPR:
                clean.remove(child)
        body = etree.tostring(clean, method="c14n").decode("utf-8")
    return f"{wrapper_index}|{body}"


class _Walk:
    """A paragraph as pieces (for `runs.encode`) and the children it keeps."""

    def __init__(self, p):
        self.pieces: list[Piece] = []
        self.kept: list = []
        self._wrappers = 0
        self._field: list | None = None
        self._depth = 0
        self._walk(p, None, 0)
        if self._field:  # a field never closed: keep it whole all the same
            self.pieces.append(Piece("object", handle=("elements", self._field, None)))

    def _field_depth(self, run) -> int:
        d = 0
        for fc in run.iter(W + "fldChar"):
            t = fc.get(W + "fldCharType")
            d += 1 if t == "begin" else -1 if t == "end" else 0
        return d

    def _walk(self, container, wrapper, index: int) -> None:
        for el in container:
            tag = el.tag
            if self._field is not None:
                self._field.append(el)
                if tag == W + "r":
                    self._depth += self._field_depth(el)
                if self._depth <= 0:
                    self.pieces.append(Piece("object", handle=("elements", self._field, wrapper)))
                    self._field = None
                continue
            if tag == W + "r":
                if any(fc.get(W + "fldCharType") == "begin" for fc in el.iter(W + "fldChar")):
                    self._field, self._depth = [el], self._field_depth(el)
                    if self._depth <= 0:
                        self.pieces.append(Piece("object", handle=("elements", self._field, wrapper)))
                        self._field = None
                    continue
                self._run(el, wrapper, index)
            elif tag in _WRAPPERS:
                self._wrappers += 1
                self._walk(el, el, self._wrappers)
            elif tag in _WHOLE:
                self.pieces.append(Piece("object", handle=("elements", [el], wrapper)))
            elif tag != W + "pPr":
                self.kept.append(el)

    def _run(self, run, wrapper, index: int) -> None:
        key = _key(run.find(W + "rPr"), index)
        for child in run:
            tag = child.tag
            if tag == W + "rPr" or tag == W + "lastRenderedPageBreak":
                continue  # Word recomputes rendered page breaks
            if tag == W + "t":
                self.pieces.append(Piece("text", child.text or "", key, handle=(run, wrapper)))
            elif tag in (W + "br", W + "cr") and child.get(W + "type") in (None, "textWrapping"):
                self.pieces.append(Piece("break", handle=(run, child, wrapper)))
            else:
                self.pieces.append(Piece("object", handle=("child", run, child, wrapper)))


def _shell(run):
    """A copy of `run` holding nothing but its formatting."""
    r = copy.deepcopy(run)
    for child in list(r):
        if child.tag != W + "rPr":
            r.remove(child)
    return r


def _text_run(template, text: str):
    r = _shell(template)
    for i, part in enumerate(_NEWLINES.split(text)):
        if i:
            r.append(OxmlElement("w:br"))
        if part:
            t = OxmlElement("w:t")
            t.set(XML_SPACE, "preserve")
            t.text = part
            r.append(t)
    return r


def _rebuild_paragraph(p, walk: _Walk, layout, outs, placeholders) -> None:
    new: list = []
    open_src, open_new = None, None

    def emit(el, wrapper) -> None:
        nonlocal open_src, open_new
        if wrapper is None:
            open_src = None
            new.append(el)
            return
        if wrapper is not open_src:
            open_new = copy.deepcopy(wrapper)
            for child in list(open_new):
                open_new.remove(child)
            open_src = wrapper
            new.append(open_new)
        open_new.append(el)

    breaks = iter(layout.breaks)
    base_run, base_wrapper = layout.spans[layout.base][0].handle if layout.base >= 0 else (None, None)
    for o in outs:
        if o.kind == "text":
            run, wrapper = layout.spans[o.span][0].handle
            emit(_text_run(run, restore(o.text, placeholders)), wrapper)
        elif o.kind == "object":
            h = o.piece.handle
            if h[0] == "elements":
                for el in h[1]:
                    emit(el, h[2])  # moved, not copied: a text box inside stays live
            else:
                _, run, child, wrapper = h
                r = _shell(run)
                r.append(child)
                emit(r, wrapper)
        else:
            src = next(breaks, None)
            if src is not None:
                run, child, wrapper = src.handle
                r = _shell(run)
                r.append(child)
                emit(r, wrapper)
            elif base_run is not None:
                r = _shell(base_run)
                r.append(OxmlElement("w:br"))
                emit(r, None)

    ppr = p.find(W + "pPr")
    leading = [k for k in walk.kept if k.tag.endswith("Start")]
    trailing = [k for k in walk.kept if not k.tag.endswith("Start") and k.tag != W + "proofErr"]
    for child in list(p):
        p.remove(child)
    for el in ([ppr] if ppr is not None else []) + leading + new + trailing:
        p.append(el)


def _story_parts(document) -> list[tuple[str, object]]:
    """(name, part) for the main story and every header/footer part."""
    out = [("document", document.part)]
    seen = set()
    for rel in document.part.rels.values():
        if rel.is_external or rel.reltype not in (RT.HEADER, RT.FOOTER):
            continue
        part = rel.target_part
        name = str(part.partname).rsplit("/", 1)[-1].rsplit(".", 1)[0]
        if name not in seen:
            seen.add(name)
            out.append((name, part))
    return [out[0]] + sorted(out[1:], key=lambda x: x[0])


def _parts(document) -> list[tuple[str, object]]:
    """(name, root element) for the main story and every header/footer part."""
    return [(name, part.element) for name, part in _story_parts(document)]


def paragraphs(document):
    """(segment id, section index, w:p) for every paragraph the adapter reads,
    in a fixed order — extraction and rebuilding walk exactly this."""
    for name, root in _parts(document):
        section = 0
        for n, p in enumerate(list(root.iter(W + "p"))):
            if any(a.tag in (W + "del", W + "moveFrom") for a in p.iterancestors()):
                continue
            yield f"docx:{name}.p{n}", section, p
            if name == "document" and p.find(f"{W}pPr/{W}sectPr") is not None:
                section += 1


class DocxAdapter:
    def extract(self, src: str):
        document = python_docx.Document(src)
        segs = []
        for seg_id, section, p in paragraphs(document):
            walk = _Walk(p)
            if not any(pc.kind == "text" and pc.text.strip() for pc in walk.pieces):
                continue
            tagged, _ = encode(walk.pieces)
            segs.append(blank_segment(seg_id, tagged, page=section))
        return segs

    def rebuild(self, src, segments, out, lang):
        document = python_docx.Document(src)
        by_id = {s.id: s for s in segments}
        issues: list[Issue] = []
        written = []
        roots = {id(part.element): part for _, part in _story_parts(document)}
        touched = {}
        for seg_id, _, p in list(paragraphs(document)):
            seg = by_id.get(seg_id)
            if not writable(seg):
                continue
            walk = _Walk(p)
            _, layout = encode(walk.pieces)
            outs = decode(seg.target, layout)
            if outs is None:
                issues.append(reject_tags(seg, seg_id))
                continue
            _rebuild_paragraph(p, walk, layout, outs, seg.placeholders)
            written.append(p)
            part = roots.get(id(p.getroottree().getroot()))
            if part is not None:
                touched[member(part.partname)] = part
        if lang.direction == "rtl" and written:
            from pagebirdy.office import rtl
            rtl.apply_docx(written, lang)
        # Only the story parts that changed are re-serialised; styles,
        # numbering, settings, media and relationships stay byte for byte.
        write_patched(src, out, {name: part_bytes(part.element) for name, part in touched.items()})
        return issues

    def validate(self, src, out):
        try:
            a, b = python_docx.Document(src), python_docx.Document(out)
        except Exception as e:  # noqa: BLE001 — any failure to open is the finding
            return [Issue("error", "unreadable", "output", f"output does not open: {e}")]
        issues = []
        sa, sb = _census(a), _census(b)
        for what in sa:
            if sa[what] != sb[what]:
                issues.append(Issue("error", "structure", what,
                                    f"{sa[what]} in source, {sb[what]} in output"))
        for name, root in _parts(b):
            if "⟦" in etree.tostring(root, encoding="unicode"):
                issues.append(Issue("error", "tokens", name, "a placeholder token reached the file"))
        return issues

    def units(self, src):
        d = python_docx.Document(src)
        root = d.part.element
        return {"sections": len(d.sections),
                "paragraphs": sum(1 for _ in root.iter(W + "p")),
                "tables": sum(1 for _ in root.iter(W + "tbl"))}


def _census(document) -> dict:
    counts = {"parts": 0, "paragraphs": 0, "tables": 0, "drawings": 0,
              "page breaks": 0, "sections": 0, "paragraph styles": ()}
    styles = set()
    for _, root in _parts(document):
        counts["parts"] += 1
        counts["paragraphs"] += sum(1 for _ in root.iter(W + "p"))
        counts["tables"] += sum(1 for _ in root.iter(W + "tbl"))
        counts["drawings"] += sum(1 for _ in root.iter(W + "drawing"))
        counts["page breaks"] += sum(1 for br in root.iter(W + "br") if br.get(W + "type") == "page")
        counts["sections"] += sum(1 for _ in root.iter(W + "sectPr"))
        styles |= {s.get(W + "val") for s in root.iter(W + "pStyle")}
    counts["paragraph styles"] = tuple(sorted(styles))
    return counts


ADAPTER = DocxAdapter()
FORMAT = Format(".docx", "docx", "DOCX",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                100 * MB, "word/document.xml", "pagebirdy.office.docx")
