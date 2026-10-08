"""Text in pictures on PPTX slides: OCR'd once, translated with the deck,
reviewed with the deck, redrawn in place. No test calls a real OCR service:
`office.images._ocr` is replaced by the synthetic scene's own regions."""

import copy
import io
import os
import zipfile

import numpy as np
import pytest
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches

from pagebirdy.image.pipeline import analyze_raster
from pagebirdy.office import images, pipeline
from pagebirdy.review.store import ReviewStore
from tests.image_fixtures import DictEngine, Scene, TextItem, fake_ocr, render_scene
from tests.office_fixtures import LANG_ES, png_bytes

TABLE = {"SALE TODAY": "REBAJAS HOY"}


def _scene_png(fmt="PNG", x=40):
    img, regions = render_scene(Scene(items=[TextItem(["SALE TODAY"], x, 140, size=40)]))
    buf = io.BytesIO()
    img.save(buf, fmt, **({"quality": 95} if fmt == "JPEG" else {}))
    return buf.getvalue(), regions


@pytest.fixture
def ocr(monkeypatch):
    """Turns picture OCR on and counts calls; `ocr.regions` is what it reads."""
    monkeypatch.setenv("BABEL_OFFICE_IMAGE_OCR", "1")
    monkeypatch.setenv("BABEL_IMAGE_OCR_ENGINE", "rapidocr")

    class Fake:
        calls = 0
        regions: list = []

        def __call__(self, work):
            Fake.calls += 1
            return analyze_raster(work, LANG_ES, ocr=fake_ocr(self.regions))

    fake = Fake()
    monkeypatch.setattr(images, "_ocr", fake)
    return fake


def _deck(path, blob, slides=1, crop_right=0.0):
    prs = Presentation()
    for _ in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic = slide.shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(1), width=Inches(5))
        pic.crop_right = crop_right
    prs.save(path)


def _run(tmp_path, lang="es", table=TABLE):
    store = ReviewStore(str(tmp_path / "r.db"))
    job_id = store.create_pending_job(str(tmp_path / "in.pptx"), meta={"target_lang": lang})
    store.close()
    return pipeline.translate_document(
        str(tmp_path / "in.pptx"), str(tmp_path / "out"), target_lang=lang,
        review_db=str(tmp_path / "r.db"), job_id=job_id,
        engines=(DictEngine(table), None))


def _members(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def _media(path):
    return {n: b for n, b in _members(path).items() if n.startswith("ppt/media/")}


def test_picture_text_is_translated_and_redrawn_in_place(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path)
    assert rep["errors"] == []
    assert rep["images_translated"] == 1 and rep["units"]["images"] == 1
    src, out = _members(tmp_path / "in.pptx"), _members(rep["output"])
    changed = {n for n in src if src[n] != out.get(n)}
    assert changed == {"ppt/media/image1.png"}
    new = Image.open(io.BytesIO(out["ppt/media/image1.png"]))
    old = Image.open(io.BytesIO(src["ppt/media/image1.png"]))
    assert (new.format, new.size) == (old.format, old.size)
    assert not np.array_equal(np.asarray(new), np.asarray(old))
    store = ReviewStore(str(tmp_path / "r.db"))
    rows = store.get_segments(rep["job_id"])
    store.close()
    assert [(r["seg_id"], r["target"]) for r in rows] == [("pptx:img.image1-png.r1", "REBAJAS HOY")]


def test_one_picture_on_many_slides_is_read_once(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob, slides=3)
    rep = _run(tmp_path)
    assert ocr.calls == 1
    assert rep["segments_total"] == 1 and rep["images_translated"] == 1


def test_jpeg_stays_jpeg(tmp_path, ocr):
    blob, ocr.regions = _scene_png("JPEG")
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path)
    assert rep["errors"] == []
    (name, data), = _media(rep["output"]).items()
    assert name.endswith((".jpg", ".jpeg")) and Image.open(io.BytesIO(data)).format == "JPEG"
    assert data != _media(tmp_path / "in.pptx")[name]


def test_text_cropped_out_of_view_is_not_translated(tmp_path, ocr):
    blob, ocr.regions = _scene_png(x=300)   # text on the right of a 480px picture
    _deck(str(tmp_path / "in.pptx"), blob, crop_right=0.5)
    rep = _run(tmp_path)
    assert rep["segments_total"] == 0
    assert _media(rep["output"]) == _media(tmp_path / "in.pptx")


def test_tiny_picture_is_never_read(tmp_path, ocr):
    _deck(str(tmp_path / "in.pptx"), png_bytes())
    rep = _run(tmp_path)
    assert ocr.calls == 0 and rep["segments_total"] == 0
    assert [w for w in rep["warnings"] if w["code"] == "image"] == []


def test_switched_off_leaves_pictures_alone(tmp_path, ocr, monkeypatch):
    monkeypatch.setenv("BABEL_OFFICE_IMAGE_OCR", "0")
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path)
    assert ocr.calls == 0 and rep["segments_total"] == 0
    assert "images" not in rep["units"]
    assert _media(rep["output"]) == _media(tmp_path / "in.pptx")


def test_rtl_never_touches_the_picture_shape(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path, lang="ar", table={"SALE TODAY": "تخفيضات اليوم"})
    assert rep["errors"] == [] and rep["images_translated"] == 1
    src, out = _members(tmp_path / "in.pptx"), _members(rep["output"])
    assert src["ppt/slides/slide1.xml"] == out["ppt/slides/slide1.xml"]


def _edit_and_rebuild(tmp_path, rep, target):
    store = ReviewStore(str(tmp_path / "r.db"))
    store.update_segment(rep["job_id"], "pptx:img.image1-png.r1", target, approve=True)
    job = store.get_job(rep["job_id"])
    store.close()
    return pipeline.rebuild_from_review(job, str(tmp_path / "r.db"))


def test_review_edit_is_redrawn_without_ocr(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path)
    first = _media(rep["output"])
    calls = ocr.calls
    out = _edit_and_rebuild(tmp_path, rep, "OFERTA")
    assert ocr.calls == calls
    rebuilt = _media(out)
    assert rebuilt != first and rebuilt != _media(tmp_path / "in.pptx")


def test_without_the_ocr_cache_a_rebuild_keeps_the_picture(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    _deck(str(tmp_path / "in.pptx"), blob)
    rep = _run(tmp_path)
    os.remove(images.cache_path(str(tmp_path / "in.pptx")))
    out = _edit_and_rebuild(tmp_path, rep, "OFERTA")
    assert ocr.calls == 1
    assert _media(out) == _media(tmp_path / "in.pptx")


def test_pictures_on_layouts_are_reported_not_translated(tmp_path, ocr):
    blob, ocr.regions = _scene_png()
    other, _ = _scene_png(x=60)
    path = str(tmp_path / "in.pptx")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    pic = slide.shapes.add_picture(io.BytesIO(other), Inches(1), Inches(1), width=Inches(5))
    layout = prs.slide_layouts[6]
    rid = layout.part.relate_to(slide.part.related_part(pic._element.blipFill.blip.rEmbed),
                                RT.IMAGE)
    el = copy.deepcopy(pic._element)
    el.blipFill.blip.rEmbed = rid
    layout.shapes._spTree.append(el)
    pic._element.getparent().remove(pic._element)
    slide.shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(4), width=Inches(5))
    prs.save(path)
    rep = _run(tmp_path)
    assert rep["images_translated"] == 1
    assert any(w["where"] == "layouts/masters" for w in rep["warnings"])
