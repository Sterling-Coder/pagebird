"""Right-to-left output for DOCX and PPTX.

The alignment rules here were measured, not assumed (Word and PowerPoint via
COM, exported to PDF, text position read back):
  * Word: in a `w:bidi` paragraph `w:jc` left/right are the leading/trailing
    edge — `left` renders on the right. Setting `w:bidi` alone mirrors it.
  * PowerPoint: `a:pPr/@algn` is physical — `rtl="1"` with no `algn` still
    renders on the left. Alignment has to be resolved and swapped.
"""

import zipfile

import docx as python_docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation

from pagebirdy import languages
from pagebirdy.office.docx import ADAPTER as DOCX
from pagebirdy.office.pipeline import prepare
from pagebirdy.office.pptx import ADAPTER as PPTX
from tests.office_fixtures import LANG_ES, make_docx, make_pptx, translate_all

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def AR(s):
    return s.replace("Revenue", "الإيرادات").replace("Total", "إجمالي").replace("Click", "انقر")


def _out(tmp_path, adapter, maker, lang, ext, fn=AR):
    src = str(tmp_path / f"in.{ext}")
    maker(src)
    segs = adapter.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, fn)
    out = str(tmp_path / f"out.{ext}")
    issues = adapter.rebuild(src, segs, out, lang) + adapter.validate(src, out)
    return src, out, issues


@pytest.mark.parametrize("code", ["ar", "he", "fa", "ur"])
def test_docx_rtl_attributes(tmp_path, code):
    _, out, issues = _out(tmp_path, DOCX, make_docx, languages.get(code), "docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "<w:bidi/>" in xml and "<w:rtl/>" in xml and "<w:bidiVisual/>" in xml
    assert f'w:bidi="{code}' in xml and 'w:cs="' in xml
    assert "100000" in xml and "help@acme.com" in xml
    assert [i for i in issues if i.level == "error"] == []
    assert python_docx.Document(out).paragraphs  # still opens


def test_docx_ltr_has_no_rtl_attributes(tmp_path):
    _, out, _ = _out(tmp_path, DOCX, make_docx, LANG_ES, "docx")
    xml = zipfile.ZipFile(out).read("word/document.xml").decode("utf-8")
    assert "w:bidi" not in xml and "<w:rtl/>" not in xml


def test_docx_alignment_is_left_to_bidi(tmp_path):
    """jc is relative in a bidi paragraph, so it is written back unchanged."""
    src = str(tmp_path / "a.docx")
    d = python_docx.Document()
    d.add_paragraph("Total left").alignment = WD_ALIGN_PARAGRAPH.LEFT
    d.add_paragraph("Total right").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("Total centre").alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(src)
    segs = DOCX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, AR)
    out = str(tmp_path / "b.docx")
    DOCX.rebuild(src, segs, out, languages.get("ar"))
    assert [p.alignment for p in python_docx.Document(out).paragraphs] == [
        WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER]


def test_docx_rpr_children_stay_in_schema_order(tmp_path):
    """Word rejects an rPr whose children are out of CT_RPr order."""
    _, out, _ = _out(tmp_path, DOCX, make_docx, languages.get("ar"), "docx")
    root = python_docx.Document(out).part.element
    order = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
             "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
             "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
             "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
             "specVanish", "oMath"]
    for rpr in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr"):
        names = [c.tag.split("}")[1] for c in rpr if c.tag.split("}")[1] in order]
        assert names == sorted(names, key=order.index), names


def _algn(out):
    slide = Presentation(out).slides[1]
    res = {}
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text:
            ppr = sh.text_frame.paragraphs[0]._p.find(A + "pPr")
            res[sh.text_frame.text] = (ppr.get("rtl") if ppr is not None else None,
                                       ppr.get("algn") if ppr is not None else None)
    return res


def test_pptx_rtl_attributes_alignment_and_geometry(tmp_path):
    src, out, issues = _out(tmp_path, PPTX, make_pptx, languages.get("ar"), "pptx")
    xml = zipfile.ZipFile(out).read("ppt/slides/slide2.xml").decode("utf-8")
    assert 'rtl="1"' in xml and 'lang="ar-SA"' in xml and "<a:cs " in xml
    # no explicit or inherited alignment -> physical left -> mirrored to right
    assert _algn(out)["الإيرادات grew 12%"] == ("1", "r")
    assert [i for i in issues if i.level == "error"] == []


def test_pptx_explicit_alignment_swaps_and_centre_stays(tmp_path):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    src = str(tmp_path / "a.pptx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (text, al) in enumerate([("Total left", PP_ALIGN.LEFT), ("Total right", PP_ALIGN.RIGHT),
                                    ("Total centre", PP_ALIGN.CENTER)]):
        tb = s.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(5), Inches(0.8))
        tb.text_frame.text = text
        tb.text_frame.paragraphs[0].alignment = al
    prs.save(src)
    segs = PPTX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, AR)
    out = str(tmp_path / "b.pptx")
    PPTX.rebuild(src, segs, out, languages.get("ar"))
    got = [p.text_frame.paragraphs[0]._p.find(A + "pPr").get("algn") for p in Presentation(out).slides[0].shapes]
    assert got == ["r", "l", "ctr"]


def test_pptx_inherited_title_alignment_is_respected(tmp_path):
    """The default template's title is centred by its layout; it must stay centred."""
    src = str(tmp_path / "in.pptx")
    make_pptx(src)
    segs = PPTX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, AR)
    out = str(tmp_path / "out.pptx")
    PPTX.rebuild(src, segs, out, languages.get("ar"))
    title = Presentation(out).slides[0].shapes.title.text_frame.paragraphs[0]._p.find(A + "pPr")
    assert title.get("rtl") == "1" and title.get("algn") in (None, "ctr")


def test_pptx_ltr_unchanged_direction(tmp_path):
    _, out, _ = _out(tmp_path, PPTX, make_pptx, LANG_ES, "pptx")
    xml = zipfile.ZipFile(out).read("ppt/slides/slide2.xml").decode("utf-8")
    assert 'rtl="1"' not in xml and "<a:cs " not in xml
