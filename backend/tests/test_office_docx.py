import zipfile

import docx
import pytest

from pagebirdy.office import formats
from pagebirdy.office.docx import ADAPTER as DOCX
from pagebirdy.office.pipeline import prepare
from tests.office_fixtures import LANG_ES, make_docx, translate_all


def _build(tmp_path, fn, mutate=None):
    src = str(tmp_path / "in.docx"); make_docx(src)
    segs = DOCX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, fn)
    if mutate:
        mutate(todo)
    out = str(tmp_path / "out.docx")
    issues = DOCX.rebuild(src, segs, out, LANG_ES)
    return src, out, segs, issues + DOCX.validate(src, out)


def _texts(path):
    return [p.text for p in docx.Document(path).paragraphs]


def test_extracts_body_table_header_footer(tmp_path):
    src = str(tmp_path / "in.docx"); make_docx(src)
    segs = DOCX.extract(src)
    texts = [s.source for s in segs]
    assert "Click ⟦r1⟧Save⟦/r1⟧ to continue." in texts
    assert "Total Revenue" in texts and "Confidential" in texts and "Page footer" in texts
    assert "Name:⟦x0⟧value" in texts
    assert "Read the ⟦r1⟧user guide⟦/r1⟧ first." in texts
    ids = [s.id for s in segs]
    assert any(i.startswith("docx:header") for i in ids) and any(i.startswith("docx:footer") for i in ids)
    assert len(ids) == len(set(ids))


def test_identity_roundtrip_keeps_everything(tmp_path):
    src, out, _, issues = _build(tmp_path, lambda s: s)
    assert [i for i in issues if i.level == "error"] == []
    a, b = docx.Document(src), docx.Document(out)
    assert _texts(src) == _texts(out)
    assert [p.style.name for p in a.paragraphs] == [p.style.name for p in b.paragraphs]
    assert len(b.inline_shapes) == 1 and len(b.tables) == 1
    assert b.sections[0].header.paragraphs[0].text == "Confidential"
    xml = zipfile.ZipFile(out).read("word/document.xml")
    assert b'w:type="page"' in xml and "⟦".encode() not in xml


def test_bold_follows_the_translated_word(tmp_path):
    def es(s):
        return {"Click ⟦r1⟧Save⟦/r1⟧ to continue.": "Pulsa ⟦r1⟧Guardar⟦/r1⟧ para seguir."}.get(s, s)
    _, out, _, _ = _build(tmp_path, es)
    p = [p for p in docx.Document(out).paragraphs if p.text.startswith("Pulsa")][0]
    assert [(r.text, bool(r.bold)) for r in p.runs if r.text] == [
        ("Pulsa ", False), ("Guardar", True), (" para seguir.", False)]
    assert str([r for r in p.runs if r.text == "Guardar"][0].font.color.rgb) == "C00000"


def test_hyperlink_stays_a_link_around_the_translated_words(tmp_path):
    def es(s):
        return {"Read the ⟦r1⟧user guide⟦/r1⟧ first.": "Lee primero la ⟦r1⟧guía⟦/r1⟧."}.get(s, s)
    _, out, _, _ = _build(tmp_path, es)
    p = [p for p in docx.Document(out).paragraphs if p.text.startswith("Lee")][0]
    links = p._p.xpath("./w:hyperlink")
    assert len(links) == 1 and links[0].xpath("string(.//w:t)") == "guía"
    assert p.text == "Lee primero la guía."


def test_protected_literals_unchanged(tmp_path):
    _, out, _, _ = _build(tmp_path, lambda s: s.replace("Email", "Correo"))
    assert "Correo help@acme.com or see https://acme.com/help." in _texts(out)
    assert docx.Document(out).tables[0].cell(0, 1).text == "100000"


def test_reordered_tags_keep_source_and_flag(tmp_path):
    def broken(todo):
        for s in todo:
            if s.source.startswith("Click"):
                s.target = "Pulsa ⟦/r1⟧Guardar⟦r1⟧ ya"
    _, out, segs, issues = _build(tmp_path, lambda s: s, mutate=broken)
    assert any(i.code == "tags" for i in issues)
    assert "Click Save to continue." in _texts(out)
    assert [s.status for s in segs if s.source.startswith("Click")] == ["needs_human"]


def test_tab_object_survives_translation(tmp_path):
    _, out, _, _ = _build(tmp_path, lambda s: s.replace("Name:", "Nombre:"))
    p = [p for p in docx.Document(out).paragraphs if p.text.startswith("Nombre:")][0]
    assert p._p.xpath(".//w:tab")


def test_needs_human_paragraph_is_left_as_source(tmp_path):
    def hold(todo):
        for s in todo:
            if s.source == "Total Revenue":
                s.status = "needs_human"
    _, out, _, _ = _build(tmp_path, lambda s: "X" + s, mutate=hold)
    assert docx.Document(out).tables[0].cell(0, 0).text == "Total Revenue"


def test_docx_content_check(tmp_path):
    import zipfile as zf
    fmt = formats.lookup("a.docx")
    assert fmt is not None
    good = tmp_path / "a.docx"; make_docx(str(good))
    formats.check_content(fmt, str(good))
    fake = tmp_path / "b.docx"
    with zf.ZipFile(fake, "w") as z:
        z.writestr("ppt/presentation.xml", "<x/>")
    png = tmp_path / "c.docx"; png.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 64)
    trunc = tmp_path / "d.docx"; trunc.write_bytes(good.read_bytes()[:200])
    for bad in (fake, png, trunc):
        with pytest.raises(formats.UnsupportedFile):
            formats.check_content(fmt, str(bad))


def test_only_edited_parts_change(tmp_path):
    src, out, _, _ = _build(tmp_path, lambda s: s.replace("Confidential", "Confidencial"))
    za, zb = zipfile.ZipFile(src), zipfile.ZipFile(out)
    assert za.namelist() == zb.namelist()
    changed = sorted(n for n in za.namelist() if za.read(n) != zb.read(n))
    assert all(n == "word/document.xml" or n.startswith(("word/header", "word/footer")) for n in changed)
    assert "word/styles.xml" not in changed and "word/numbering.xml" not in changed
