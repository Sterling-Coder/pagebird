import zipfile

import pytest
from pptx import Presentation

from pagebirdy.office import formats
from pagebirdy.office.pipeline import prepare
from pagebirdy.office.pptx import ADAPTER as PPTX
from tests.office_fixtures import LANG_ES, make_pptx, translate_all


def _geometry(path):
    prs = Presentation(path)
    out = [(prs.slide_width, prs.slide_height, len(prs.slides))]
    for slide in prs.slides:
        for sh in slide.shapes:
            out.append((sh.shape_id, sh.left, sh.top, sh.width, sh.height))
    return out


def _build(tmp_path, fn, mutate=None):
    src = str(tmp_path / "in.pptx"); make_pptx(src)
    segs = PPTX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, fn)
    if mutate:
        mutate(todo)
    out = str(tmp_path / "out.pptx")
    issues = PPTX.rebuild(src, segs, out, LANG_ES) + PPTX.validate(src, out)
    return src, out, segs, issues


def test_extracts_titles_boxes_tables_groups_not_notes(tmp_path):
    src = str(tmp_path / "in.pptx"); make_pptx(src)
    segs = PPTX.extract(src)
    texts = [s.source for s in segs]
    for t in ("Quarterly Results", "Prepared by Finance", "Region", "North", "Grouped label"):
        assert t in texts
    assert "Revenue grew ⟦r1⟧12%⟦/r1⟧" in texts
    assert "Speaker says hello" not in texts
    assert len({s.id for s in segs}) == len(segs)
    assert {s.page for s in segs} == {1, 2}


def test_translation_lands_and_geometry_is_identical(tmp_path):
    src, out, _, issues = _build(tmp_path, lambda s: s.replace("Region", "Región"))
    assert [i for i in issues if i.level == "error"] == []
    assert _geometry(src) == _geometry(out)
    tbl = [sh for sh in Presentation(out).slides[1].shapes if sh.has_table][0].table
    assert tbl.cell(0, 0).text == "Región" and tbl.cell(1, 1).text == "4500"
    assert Presentation(out).slides[1].notes_slide.notes_text_frame.text == "Speaker says hello"


def test_bold_run_keeps_its_formatting(tmp_path):
    _, out, _, _ = _build(tmp_path, lambda s: s.replace("Revenue grew ", "Los ingresos subieron "))
    box = [sh for sh in Presentation(out).slides[1].shapes if sh.has_text_frame
           and sh.text_frame.text.startswith("Los")][0]
    runs = box.text_frame.paragraphs[0].runs
    assert [(r.text, bool(r.font.bold)) for r in runs] == [("Los ingresos subieron ", False), ("12%", True)]


def test_group_text_and_master_untouched(tmp_path):
    src, out, _, _ = _build(tmp_path, lambda s: s.replace("Grouped label", "Etiqueta"))
    za, zb = zipfile.ZipFile(src), zipfile.ZipFile(out)
    for name in za.namelist():
        if "slideMaster" in name or "slideLayout" in name or "theme" in name or "notesSlide" in name:
            assert za.read(name) == zb.read(name), name
    assert "Etiqueta" in zb.read("ppt/slides/slide2.xml").decode("utf-8")


def test_reordered_tags_keep_source(tmp_path):
    def broken(todo):
        for s in todo:
            if "⟦r1⟧" in s.source:
                s.target = "⟦/r1⟧x⟦r1⟧"
    _, out, segs, issues = _build(tmp_path, lambda s: s, mutate=broken)
    assert any(i.code == "tags" for i in issues)
    texts = [sh.text_frame.text for sh in Presentation(out).slides[1].shapes if sh.has_text_frame]
    assert "Revenue grew 12%" in texts
    assert [s.status for s in segs if "⟦r1⟧" in s.source] == ["needs_human"]


def test_pptx_content_check(tmp_path):
    fmt = formats.lookup("deck.pptx")
    assert fmt is not None
    good = tmp_path / "a.pptx"; make_pptx(str(good))
    formats.check_content(fmt, str(good))
    fake = tmp_path / "b.pptx"
    with zipfile.ZipFile(fake, "w") as z:
        z.writestr("word/document.xml", "<x/>")
    for bad in (fake,):
        with pytest.raises(formats.UnsupportedFile):
            formats.check_content(fmt, str(bad))
