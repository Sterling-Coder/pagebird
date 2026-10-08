from pptx import Presentation
from pptx.util import Inches, Pt

from pagebirdy.office.pipeline import prepare
from pagebirdy.office.pptx import ADAPTER as PPTX
from tests.office_fixtures import LANG_ES, translate_all


def _deck(path, text, w=3.0, h=0.6, size=18, autofit=None):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(size)
    if autofit == "norm":
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    prs.save(path)


def _run(tmp_path, text, fn, **kw):
    src = str(tmp_path / "in.pptx")
    _deck(src, text, **kw)
    segs = PPTX.extract(src)
    todo, _ = prepare(segs)
    translate_all(todo, fn)
    out = str(tmp_path / "out.pptx")
    issues = PPTX.rebuild(src, segs, out, LANG_ES)
    sh = Presentation(out).slides[0].shapes[0]
    sizes = [r.font.size.pt for r in sh.text_frame.paragraphs[0].runs]
    return sh, sizes, issues, out


def test_fitting_text_keeps_source_size(tmp_path):
    _, sizes, issues, out = _run(tmp_path, "Short", lambda s: "Corto")
    assert sizes == [18.0] and issues == []
    assert "lnSpc" not in Presentation(out).slides[0].shapes[0].text_frame._txBody.xml


def test_longer_text_shrinks_within_floor(tmp_path):
    sh, sizes, issues, _ = _run(tmp_path, "Quarterly revenue summary",
                                lambda s: "Resumen trimestral de ingresos y gastos operativos")
    assert 18 * 0.7 <= sizes[0] < 18
    assert sizes[0] == int(sizes[0])  # whole points
    assert sh.width == Inches(3) and sh.height == Inches(0.6)
    assert not [i for i in issues if i.code == "overflow"]


def test_line_spacing_is_tried_before_shrinking(tmp_path):
    # Two lines that overflow at 100% spacing but fit at 90%.
    sh, sizes, issues, _ = _run(tmp_path, "x", lambda s: "word " * 8, w=3.0, h=0.66)
    assert sizes == [18.0] and issues == []
    assert 'spcPct val="90000"' in sh.text_frame._txBody.xml


def test_impossible_fit_warns_with_slide_number(tmp_path):
    _, sizes, issues, _ = _run(tmp_path, "Hi", lambda s: "palabra " * 80, w=2.0, h=0.4)
    assert sizes[0] == 12.6  # max(18 * 0.7, 8)
    assert [i.where for i in issues if i.code == "overflow"] == ["slide 1"]


def test_norm_autofit_frames_left_to_powerpoint(tmp_path):
    _, sizes, issues, _ = _run(tmp_path, "Hi", lambda s: "palabra " * 80, autofit="norm")
    assert sizes == [18.0] and issues == []


def test_inherited_title_size_is_resolved(tmp_path):
    from pagebirdy.office.pptx_fit import resolved_size
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Title"
    run = s.shapes.title.text_frame.paragraphs[0].runs[0]._r
    assert resolved_size(s, s.shapes.title, run) == 44.0  # default template title size


def test_source_that_already_overflows_is_not_shrunk_when_unchanged(tmp_path):
    """The room a frame has is at least what the source needed: an estimator
    that says the English overflows must not shrink an identical translation."""
    text = "palabra " * 30
    _, sizes, issues, _ = _run(tmp_path, text, lambda s: s, w=2.0, h=0.4)
    assert sizes == [18.0] and issues == []
