"""Fixture builders shared by the office tests. Nothing binary is committed:
every document is built here, in the test that needs it."""

import re

from pagebirdy import languages

LANG_ES = languages.get("es")
LANG_AR = languages.get("ar")


def translate_all(segments, fn, status="translated"):
    for s in segments:
        if s.source.strip():
            s.target = fn(s.source)
            s.status = status
    return segments


def upper_outside_tokens(text: str) -> str:
    parts = re.split(r"(⟦[^⟧]*⟧)", text)
    return "".join(p if i % 2 else p.upper() for i, p in enumerate(parts))


def png_bytes() -> bytes:
    import io

    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "red").save(buf, "PNG")
    return buf.getvalue()


def _add_hyperlink(paragraph, text, url):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle"); style.set(qn("w:val"), "Hyperlink")
    rpr.append(style); run.append(rpr)
    t = OxmlElement("w:t"); t.text = text
    run.append(t); link.append(run)
    paragraph._p.append(link)


def make_docx(path):
    """A document exercising every DOCX structure the adapter must keep."""
    import io

    import docx
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    d.add_heading("Annual Report", level=1)
    p = d.add_paragraph("Click ")
    b = p.add_run("Save"); b.bold = True; b.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run(" to continue.")
    d.add_paragraph("First item", style="List Bullet")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Total Revenue"
    t.cell(0, 1).text = "100000"
    r = d.add_paragraph().add_run("Before break"); r.add_break(WD_BREAK.PAGE)
    d.add_picture(io.BytesIO(png_bytes()))
    d.add_paragraph("Email help@acme.com or see https://acme.com/help.")
    tabbed = d.add_paragraph("Name:"); tabbed.add_run().add_tab(); tabbed.add_run("value")
    linked = d.add_paragraph("Read the ")
    _add_hyperlink(linked, "user guide", "https://acme.com/guide")
    linked.add_run(" first.")
    s = d.sections[0]
    s.header.paragraphs[0].text = "Confidential"
    s.footer.paragraphs[0].text = "Page footer"
    big = d.add_paragraph().add_run("Big"); big.font.size = Pt(20); big.italic = True
    d.save(path)


def make_pptx(path):
    import io

    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quarterly Results"
    s1.placeholders[1].text = "Prepared by Finance"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    box = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.word_wrap = True
    p = box.text_frame.paragraphs[0]
    r1 = p.add_run(); r1.text = "Revenue grew "; r1.font.size = Pt(18)
    r2 = p.add_run(); r2.text = "12%"; r2.font.size = Pt(18); r2.font.bold = True
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "Region"; tbl.cell(0, 1).text = "Sales"
    tbl.cell(1, 0).text = "North"; tbl.cell(1, 1).text = "4500"
    grp = s2.shapes.add_group_shape()
    inner = grp.shapes.add_textbox(Inches(7), Inches(1), Inches(2), Inches(1))
    inner.text_frame.text = "Grouped label"
    s2.shapes.add_picture(io.BytesIO(png_bytes()), Inches(7), Inches(3))
    s2.notes_slide.notes_text_frame.text = "Speaker says hello"
    prs.save(path)


_CT = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
       '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
       '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
       '<Default Extension="xml" ContentType="application/xml"/>'
       '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
       '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
       '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
       '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
       '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
       '</Types>')
_NS = ('xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def make_xlsx(path):
    """Two sheets: text, numbers, a date, formulas (SUM, a VLOOKUP keyed on a
    string literal, a cross-sheet ref), a rich-text string, an inline string,
    a merge, a hidden column, frozen panes and a data-validation list."""
    import zipfile

    sst = [
        "Total Revenue", "Apple", "Yes", "No", "INV-2024-001", "https://acme.com",
        '<r><t xml:space="preserve">Net </t></r><r><rPr><b/><sz val="11"/></rPr><t>profit</t></r>'
        '<r><t xml:space="preserve"> margin</t></r>',
        "Region",
    ]
    si = "".join(f"<si>{s}</si>" if s.startswith("<r>") else f"<si><t>{s}</t></si>" for s in sst)
    files = {
        "[Content_Types].xml": _CT,
        "_rels/.rels": (f'{_DECL}<Relationships xmlns="{_PKG}">'
                        f'<Relationship Id="rId1" Type="{_RELS}/officeDocument" Target="xl/workbook.xml"/>'
                        '</Relationships>'),
        "xl/_rels/workbook.xml.rels": (
            f'{_DECL}<Relationships xmlns="{_PKG}">'
            f'<Relationship Id="rId1" Type="{_RELS}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{_RELS}/worksheet" Target="worksheets/sheet2.xml"/>'
            f'<Relationship Id="rId3" Type="{_RELS}/sharedStrings" Target="sharedStrings.xml"/>'
            f'<Relationship Id="rId4" Type="{_RELS}/styles" Target="styles.xml"/></Relationships>'),
        "xl/workbook.xml": (
            f'{_DECL}<workbook {_NS}><sheets><sheet name="Summary" sheetId="1" r:id="rId1"/>'
            '<sheet name="Data" sheetId="2" r:id="rId2"/></sheets>'
            '<definedNames><definedName name="Rate">Data!$B$1</definedName></definedNames></workbook>'),
        "xl/styles.xml": (
            f'{_DECL}<styleSheet {_NS}><numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy-mm-dd"/>'
            '</numFmts><fonts count="1"><font><sz val="11"/></font></fonts><fills count="1"><fill/></fills>'
            '<borders count="1"><border/></borders><cellXfs count="2"><xf/><xf numFmtId="164"/></cellXfs>'
            '</styleSheet>'),
        "xl/sharedStrings.xml": f'{_DECL}<sst {_NS} count="{len(sst)}" uniqueCount="{len(sst)}">{si}</sst>',
        "xl/worksheets/sheet1.xml": (
            f'{_DECL}<worksheet {_NS}><sheetViews><sheetView workbookViewId="0">'
            '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
            '<cols><col min="3" max="3" width="9" hidden="1" customWidth="1"/></cols><sheetData>'
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1"><v>100000</v></c><c r="C1" s="1"><v>45413</v></c></row>'
            '<row r="2"><c r="A2" t="s"><v>6</v></c><c r="B2"><f>SUM(B1:B1)</f><v>100000</v></c>'
            '<c r="C2"><f>VLOOKUP("Apple",Data!A1:B2,2,FALSE)</f><v>3</v></c></row>'
            '<row r="3"><c r="A3" t="s"><v>4</v></c><c r="B3" t="s"><v>5</v></c>'
            '<c r="C3" t="inlineStr"><is><t>Quarterly note</t></is></c><c r="D3" t="s"><v>2</v></c></row>'
            '</sheetData><mergeCells count="1"><mergeCell ref="A4:B4"/></mergeCells>'
            '<dataValidations count="1"><dataValidation type="list" sqref="D3"><formula1>"Yes,No"</formula1>'
            '</dataValidation></dataValidations></worksheet>'),
        "xl/worksheets/sheet2.xml": (
            f'{_DECL}<worksheet {_NS}><sheetData><row r="1"><c r="A1" t="s"><v>1</v></c>'
            '<c r="B1"><v>3</v></c></row><row r="2"><c r="A2" t="s"><v>7</v></c><c r="B2"><f>Summary!B1*2</f>'
            '<v>200000</v></c></row></sheetData></worksheet>'),
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
